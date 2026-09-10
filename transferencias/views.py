# -*- encoding: utf-8 -*-
import datetime
import json
import logging
from decimal import Decimal

from django.contrib.auth.decorators import user_passes_test
from django.db.models.functions import Lower
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render

from herramientas.sql_muestras_articulos import parsear_xlsx_ajuste

from .constantes import (
    CANTIDAD_COLUMNAS,
    DEPOSITOS_HABILITADOS,
    ETIQUETAS_COLUMNAS,
    GRUPOS_HABILITADOS,
    normalizar_deposito,
)
from .models import RegistroTransferenciaDeposito
from . import sql_transferencias

logger = logging.getLogger(__name__)

HOST_TANGO = 'TRANSFER-DEP'

# Historial
TAMANO_PAGINA_HISTORIAL = 50
TOPE_EXPORTACION = 20000


def usuario_puede_transferir(user):
    """Verifica si el usuario pertenece a un grupo habilitado para transferir
    stock entre depósitos."""
    return user.groups.annotate(name_lower=Lower('name')).filter(
        name_lower__in=GRUPOS_HABILITADOS
    ).exists() or user.is_superuser


def _entero_positivo(valor_crudo, etiqueta, problemas):
    """Devuelve el entero, o None acumulando el problema en la lista."""
    try:
        # Excel puede entregar '2.0' en una celda numérica.
        como_float = float(valor_crudo)
    except (ValueError, TypeError):
        problemas.append(f'la cantidad {etiqueta} "{valor_crudo}" no es un número')
        return None

    entero = int(como_float)
    if como_float != entero:
        problemas.append(f'la cantidad {etiqueta} "{valor_crudo}" debe ser un número entero')
        return None
    if entero <= 0:
        problemas.append(f'la cantidad {etiqueta} debe ser mayor a cero')
        return None
    return entero


def _validar_forma(filas_crudas):
    """Valida forma y normaliza, sin tocar la base. Devuelve (filas, errores).

    Se acumulan TODOS los errores en vez de cortar en el primero, para que el
    usuario corrija el Excel de una sola pasada.
    """
    filas = []
    errores = []

    for indice, fila in enumerate(filas_crudas, start=1):
        if len(fila) != CANTIDAD_COLUMNAS:
            errores.append(
                f'Fila {indice}: se esperaban {CANTIDAD_COLUMNAS} columnas '
                f'({", ".join(ETIQUETAS_COLUMNAS)}), se recibieron {len(fila)}.'
            )
            continue

        (dep_o_crudo, ubic_o, art_o, cant_baja_cruda,
         dep_d_crudo, ubic_d, art_d, cant_alta_cruda) = [str(v).strip() for v in fila]

        problemas = []

        dep_o = normalizar_deposito(dep_o_crudo)
        dep_d = normalizar_deposito(dep_d_crudo)

        if dep_o is None or dep_o not in DEPOSITOS_HABILITADOS:
            problemas.append(
                f'el depósito de origen "{dep_o_crudo}" no está habilitado '
                f'({", ".join(sorted(DEPOSITOS_HABILITADOS))})'
            )
        if dep_d is None or dep_d not in DEPOSITOS_HABILITADOS:
            problemas.append(
                f'el depósito de destino "{dep_d_crudo}" no está habilitado '
                f'({", ".join(sorted(DEPOSITOS_HABILITADOS))})'
            )

        if not ubic_o:
            problemas.append('la ubicación de origen es obligatoria')
        if not ubic_d:
            problemas.append('la ubicación de destino es obligatoria')
        if not art_o:
            problemas.append('el artículo de origen es obligatorio')
        if not art_d:
            problemas.append('el artículo de destino es obligatorio')

        # Recodificar en el lugar es válido (mismo depósito y ubicación con
        # código distinto). Lo único que no tiene sentido es que las tres cosas
        # sean iguales: esa fila no haría nada.
        if (dep_o and dep_d and ubic_o and ubic_d and art_o and art_d
                and dep_o == dep_d
                and ubic_o.upper() == ubic_d.upper()
                and art_o.upper() == art_d.upper()):
            problemas.append(
                'la fila no cambia nada: depósito, ubicación y artículo son iguales en origen y destino'
            )

        cant_baja = _entero_positivo(cant_baja_cruda, 'de baja', problemas)
        cant_alta = _entero_positivo(cant_alta_cruda, 'de alta', problemas)

        # Doble control de tipeo: se piden las dos cantidades justamente para
        # que un error al tipear una de ellas frene la fila en vez de mover una
        # cantidad equivocada.
        if cant_baja is not None and cant_alta is not None and cant_baja != cant_alta:
            problemas.append(
                f'la cantidad de baja ({cant_baja}) y la de alta ({cant_alta}) deben ser iguales'
            )

        if problemas:
            errores.append(f'Fila {indice}: ' + '; '.join(problemas) + '.')
            continue

        filas.append((
            dep_o, ubic_o.upper(), art_o.upper(), cant_baja,
            dep_d, ubic_d.upper(), art_d.upper(), cant_alta,
        ))

    return filas, errores


def _json_seguro(valor):
    """Convierte a algo serializable por JSONField.

    Las cantidades vuelven de SQL Server como Decimal y las fechas como
    datetime, y ninguno de los dos entra en un JSONField. Sin esto el registro
    de auditoría falla en silencio (el try/except de _auditar lo tapa, que es
    justo lo que no queremos que pase inadvertido).
    """
    if isinstance(valor, Decimal):
        entero = int(valor)
        return entero if valor == entero else float(valor)
    if isinstance(valor, (datetime.datetime, datetime.date)):
        return valor.isoformat()
    if isinstance(valor, dict):
        return {k: _json_seguro(v) for k, v in valor.items()}
    if isinstance(valor, (list, tuple)):
        return [_json_seguro(v) for v in valor]
    return valor


def _auditar(usuario, accion, lote_id, filas, filas_con_error=None, resumen=None):
    """El stock ya se movió cuando esto corre; si falla la auditoría no debe
    parecer que la operación falló (evitaría que el usuario reintente y duplique
    el movimiento)."""
    resumen = resumen or {}
    try:
        RegistroTransferenciaDeposito.objects.create(
            usuario=usuario,
            accion=accion,
            lote_id=lote_id,
            filas=_json_seguro(filas),
            filas_con_error=_json_seguro(filas_con_error) or None,
            numero_tarea=(str(resumen.get('IdTarea')) if resumen.get('IdTarea') is not None else None),
            nro_comprobante=resumen.get('NroComprobante'),
        )
    except Exception:
        logger.exception('No se pudo guardar el registro de auditoría de transferencias (%s)', accion)


@user_passes_test(usuario_puede_transferir, login_url="/login/")
def transferencias_depositos(request):
    # Las 52 ubicaciones de los depósitos habilitados van embebidas en la
    # página (~2 KB) para poblar los desplegables del modo escaneo sin un
    # round trip por cada elección de depósito. Si el WMS no responde, la
    # pantalla igual carga y el modo manual sigue funcionando.
    try:
        ubicaciones = sql_transferencias.ubicaciones_habilitadas()
    except Exception:
        logger.exception('No se pudieron cargar las ubicaciones para el modo escaneo')
        ubicaciones = []

    return render(request, 'transferencias/index.html', {
        'depositos': sorted(DEPOSITOS_HABILITADOS.items()),
        'cantidad_columnas': CANTIDAD_COLUMNAS,
        'ubicaciones_json': json.dumps(ubicaciones),
    })


@user_passes_test(usuario_puede_transferir, login_url="/login/")
def transferencias_depositos_buscar_articulo(request):
    """Resuelve un código escaneado y devuelve dónde tiene saldo.

    Es el único endpoint del modo escaneo. Solo lee: no crea lote, no toca
    stock. Los mensajes son los que va a ver el operario con el lector en la
    mano, así que cada situación se distingue explícitamente.
    """
    if request.method != 'GET':
        return JsonResponse({'error': 'Método no permitido'}, status=405)

    codigo = (request.GET.get('codigo') or '').strip()
    if not codigo:
        return JsonResponse({'errores': ['No se recibió ningún código.']}, status=400)

    try:
        datos = sql_transferencias.buscar_articulo(codigo)
    except Exception as e:
        logger.exception('Error al buscar el artículo %r', codigo)
        return JsonResponse({'errores': [f'Error al buscar el artículo: {str(e)}']}, status=500)

    if datos is None:
        return JsonResponse({
            'encontrado': False,
            'motivo': 'no_existe',
            'mensaje': f'El código "{codigo}" no existe en Tango.',
        })

    if not datos['existe_en_wms']:
        return JsonResponse({
            'encontrado': False,
            'motivo': 'sin_wms',
            'articulo': datos['articulo'],
            'mensaje': f'El artículo {datos["articulo"]} existe en Tango pero no está dado de alta en el WMS.',
        })

    if not datos['ubicaciones']:
        habilitados = ', '.join(sorted(DEPOSITOS_HABILITADOS))
        return JsonResponse({
            'encontrado': False,
            'motivo': 'sin_saldo',
            'articulo': datos['articulo'],
            'mensaje': (f'El artículo {datos["articulo"]} no tiene saldo en ninguna ubicación '
                        f'de los depósitos habilitados ({habilitados}).'),
        })

    return JsonResponse({
        'encontrado': True,
        'articulo': datos['articulo'],
        'descripcion': datos['descripcion'],
        'usa_partidas': datos['usa_partidas'],
        'ubicaciones': datos['ubicaciones'],
    })


@user_passes_test(usuario_puede_transferir, login_url="/login/")
def transferencias_depositos_subir_archivo(request):
    if request.method != 'POST':
        return JsonResponse({'error': 'Método no permitido'}, status=405)

    archivo = request.FILES.get('archivo')
    if not archivo:
        return JsonResponse({'errores': ['No se recibió ningún archivo.']}, status=400)

    if not archivo.name.lower().endswith('.xlsx'):
        return JsonResponse({'errores': ['El archivo debe tener formato .xlsx.']}, status=400)

    try:
        filas = parsear_xlsx_ajuste(archivo)
    except Exception as e:
        return JsonResponse({'errores': [f'No se pudo leer el archivo: {str(e)}']}, status=400)

    return JsonResponse({'filas': filas})


def _procesar(request, solo_validar):
    if request.method != 'POST':
        return JsonResponse({'error': 'Método no permitido'}, status=405)

    try:
        payload = json.loads(request.body)
        filas_crudas = payload.get('filas', [])
    except (json.JSONDecodeError, AttributeError):
        return JsonResponse({'error': 'JSON inválido'}, status=400)

    filas, errores = _validar_forma(filas_crudas)

    if errores:
        return JsonResponse({'errores': errores}, status=400)

    if not filas:
        return JsonResponse({'errores': ['No se recibieron filas para procesar.']}, status=400)

    lote_id = sql_transferencias.nuevo_lote_id()

    try:
        sql_transferencias.crear_lote(lote_id, request.user.username, HOST_TANGO, filas)
    except Exception as e:
        logger.exception('Error al crear el lote de transferencias')
        return JsonResponse({'errores': [f'Error al cargar los datos: {str(e)}']}, status=500)

    try:
        filas_con_error, resumen, sugerencias = sql_transferencias.procesar_lote(
            lote_id, request.user.username, HOST_TANGO, solo_validar
        )
    except Exception as e:
        logger.exception('Error técnico al procesar el lote de transferencias %s', lote_id)
        sql_transferencias.borrar_lote(lote_id)
        return JsonResponse({'errores': [f'Error al procesar la transferencia: {str(e)}']}, status=500)

    accion = (
        RegistroTransferenciaDeposito.ACCION_VALIDAR if solo_validar
        else RegistroTransferenciaDeposito.ACCION_EJECUTAR
    )
    _auditar(request.user, accion, lote_id, filas, filas_con_error, resumen)

    return JsonResponse({
        'lote_id': str(lote_id),
        'ok': bool(resumen.get('Ok')),
        'filas_con_error': filas_con_error,
        'sugerencias': sugerencias,
        'filas_procesadas': resumen.get('FilasProcesadas'),
        'cantidad_errores': resumen.get('FilasConError'),
        'numero_tarea': resumen.get('IdTarea'),
        'nro_comprobante': resumen.get('NroComprobante'),
        'mensaje': resumen.get('Mensaje'),
    })


@user_passes_test(usuario_puede_transferir, login_url="/login/")
def transferencias_depositos_validar(request):
    return _procesar(request, solo_validar=True)


@user_passes_test(usuario_puede_transferir, login_url="/login/")
def transferencias_depositos_ejecutar(request):
    return _procesar(request, solo_validar=False)


# ============================================================================
# HISTORIAL
# ============================================================================

def _filtros_historial(request):
    """Lee los filtros de la query string.

    Por defecto arranca en el primer día del mes en curso, para que la pantalla
    abra con algo útil en vez de barrer todo el historial.
    """
    hoy = datetime.date.today()
    desde = (request.GET.get('desde') or '').strip()
    hasta = (request.GET.get('hasta') or '').strip()

    # Solo se aplica el default si no vino NINGÚN filtro de fecha: si el
    # usuario puso solo "hasta", no le imponemos un "desde" que no pidió.
    if not desde and not hasta:
        desde = hoy.replace(day=1).isoformat()

    return {
        'desde': desde or None,
        'hasta': hasta or None,
        'deposito': (request.GET.get('deposito') or '').strip() or None,
        'articulo': (request.GET.get('articulo') or '').strip() or None,
        'usuario': (request.GET.get('usuario') or '').strip() or None,
    }


@user_passes_test(usuario_puede_transferir, login_url="/login/")
def transferencias_depositos_historial(request):
    """Historial de lo que pasó por la herramienta, con totales.

    Solo lee. Muestra únicamente los lotes aplicados: lo que efectivamente
    movió stock.
    """
    filtros = _filtros_historial(request)

    try:
        pagina = int(request.GET.get('page') or 1)
    except ValueError:
        pagina = 1

    try:
        totales_deposito = sql_transferencias.historial_totales_por_deposito(filtros)
        totales_articulo = sql_transferencias.historial_totales_por_articulo(filtros, limite=100)
        movimientos, total = sql_transferencias.historial_movimientos(
            filtros, pagina=pagina, tamano=TAMANO_PAGINA_HISTORIAL
        )
        error = None
    except Exception as e:
        logger.exception('Error al consultar el historial de transferencias')
        totales_deposito, totales_articulo, movimientos, total = [], [], [], 0
        error = f'No se pudo consultar el historial: {str(e)}'

    ultima = max(1, -(-total // TAMANO_PAGINA_HISTORIAL))  # techo de la división
    pagina = min(max(1, pagina), ultima)

    # La query string sin `page`, para que los links de paginado mantengan los
    # filtros. Se arma acá y no en el template, que es donde el proyecto suele
    # rearmarla a mano condición por condición.
    parametros = request.GET.copy()
    parametros.pop('page', None)
    query_filtros = parametros.urlencode()

    return render(request, 'transferencias/historial.html', {
        'filtros': filtros,
        'depositos': sorted(DEPOSITOS_HABILITADOS.items()),
        'totales_deposito': totales_deposito,
        'totales_articulo': totales_articulo,
        'movimientos': movimientos,
        'total': total,
        'pagina': pagina,
        'ultima_pagina': ultima,
        'tamano_pagina': TAMANO_PAGINA_HISTORIAL,
        'desde_registro': (pagina - 1) * TAMANO_PAGINA_HISTORIAL + 1 if total else 0,
        'hasta_registro': min(pagina * TAMANO_PAGINA_HISTORIAL, total),
        'query_filtros': query_filtros,
        'error': error,
    })


@user_passes_test(usuario_puede_transferir, login_url="/login/")
def transferencias_depositos_historial_exportar(request):
    """Baja el historial filtrado a .xlsx, con tres hojas.

    Lee los MISMOS filtros GET que la pantalla, así lo que se descarga es
    exactamente lo que se está viendo.
    """
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    filtros = _filtros_historial(request)

    try:
        movimientos = sql_transferencias.historial_movimientos_todos(filtros, tope=TOPE_EXPORTACION)
        totales_deposito = sql_transferencias.historial_totales_por_deposito(filtros)
        totales_articulo = sql_transferencias.historial_totales_por_articulo(filtros, limite=5000)
    except Exception as e:
        logger.exception('Error al exportar el historial de transferencias')
        return HttpResponse(f'No se pudo exportar el historial: {str(e)}', status=500)

    fuente_encabezado = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    relleno_encabezado = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
    centrado = Alignment(horizontal='center', vertical='center', wrap_text=True)

    libro = openpyxl.Workbook()

    def escribir(hoja, encabezados, filas):
        for columna, titulo in enumerate(encabezados, start=1):
            celda = hoja.cell(row=1, column=columna, value=titulo)
            celda.font = fuente_encabezado
            celda.fill = relleno_encabezado
            celda.alignment = centrado
        for indice, fila in enumerate(filas, start=2):
            for columna, valor in enumerate(fila, start=1):
                hoja.cell(row=indice, column=columna, value=valor)
        for columna in range(1, len(encabezados) + 1):
            ancho = len(str(encabezados[columna - 1]))
            for fila in filas:
                ancho = max(ancho, len(str(fila[columna - 1] if fila[columna - 1] is not None else '')))
            hoja.column_dimensions[get_column_letter(columna)].width = min(ancho + 3, 40)
        hoja.freeze_panes = 'A2'

    hoja = libro.active
    hoja.title = 'Movimientos'
    escribir(
        hoja,
        ['Fecha', 'Usuario', 'Tipo', 'Dep. origen', 'Ubic. origen', 'Art. origen', 'Cant. baja',
         'Dep. destino', 'Ubic. destino', 'Art. destino', 'Cant. alta', 'Comprobante', 'Tarea'],
        [
            [m['fecha'].strftime('%d/%m/%Y %H:%M') if m['fecha'] else '',
             m['usuario'],
             'Recodificación' if m['es_recodificacion'] else 'Transferencia',
             m['dep_origen'], m['ubic_origen'], m['art_origen'], m['cant_baja'],
             m['dep_destino'], m['ubic_destino'], m['art_destino'], m['cant_alta'],
             m['comprobante'], m['tarea']]
            for m in movimientos
        ],
    )

    escribir(
        libro.create_sheet('Por deposito'),
        ['Depósito', 'Nombre', 'Salidas', 'Entradas', 'Neto', 'Movimientos'],
        [[d['deposito'], d['nombre'], d['salidas'], d['entradas'], d['neto'], d['movimientos']]
         for d in totales_deposito],
    )

    escribir(
        libro.create_sheet('Por articulo'),
        ['Artículo', 'Bajas', 'Altas'],
        [[a['articulo'], a['bajas'], a['altas']] for a in totales_articulo],
    )

    respuesta = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    nombre = f"HistorialTransferencias_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    respuesta['Content-Disposition'] = f'attachment; filename="{nombre}"'
    libro.save(respuesta)
    return respuesta
