import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import os

# Crear un nuevo libro de trabajo
libro = openpyxl.Workbook()
hoja_datos = libro.active
hoja_datos.title = 'Datos'

# Estilos
example_fill = PatternFill(start_color='E7E6E6', end_color='E7E6E6', fill_type='solid')
example_font = Font(name='Arial', size=10)

thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

# Filas de ejemplo (sin encabezados: Ubicación, Artículo, Cantidad)
ejemplos = [
    ['D001', 'ART00123', 1],
    ['D001', 'ART00456', 2],
]

for fila_idx, ejemplo in enumerate(ejemplos, start=1):
    for col_idx, valor in enumerate(ejemplo, start=1):
        celda = hoja_datos.cell(row=fila_idx, column=col_idx, value=valor)
        celda.fill = example_fill
        celda.font = example_font
        celda.border = thin_border

hoja_datos.column_dimensions['A'].width = 15
hoja_datos.column_dimensions['B'].width = 15
hoja_datos.column_dimensions['C'].width = 12

# Hoja de instrucciones
hoja_instrucciones = libro.create_sheet('Instrucciones')
instrucciones = [
    'INSTRUCCIONES PARA USO DE LA PLANTILLA',
    '',
    'La hoja "Datos" NO lleva encabezados. Cada fila representa un movimiento con 3 columnas, en este orden:',
    '  1) Código de Ubicación',
    '  2) Código de Artículo',
    '  3) Cantidad (entero)',
    '',
    'Reemplazá las filas de ejemplo por tus propios datos y guardá el archivo en formato .xlsx.',
    '',
    'Para dar de alta una muestra, usá cantidad positiva (ej. 1).',
    '',
    'Si estás usando el modo Recodificación: si alguna fila tiene cantidad negativa (baja),',
    'la suma de todas las cantidades del archivo debe dar exactamente 0 (toda baja debe',
    'compensarse con alta(s) por el mismo total).',
]

for row_num, texto in enumerate(instrucciones, 1):
    celda = hoja_instrucciones.cell(row=row_num, column=1, value=texto)
    if row_num == 1:
        celda.font = Font(bold=True, size=14)

hoja_instrucciones.column_dimensions['A'].width = 95

# Guardar archivo en apps/static para que quede disponible como descarga estática
script_dir = os.path.dirname(os.path.abspath(__file__))
project_root = os.path.dirname(script_dir)
static_dir = os.path.join(project_root, 'apps', 'static')

if not os.path.exists(static_dir):
    os.makedirs(static_dir)

ruta_archivo = os.path.join(static_dir, 'PlantillaAltaMuestrasArticulos.xlsx')
libro.save(ruta_archivo)

print('Plantilla .xlsx creada exitosamente en:', ruta_archivo)
print('Columnas (sin encabezado): Código de Ubicación, Código de Artículo, Cantidad')
